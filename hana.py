"""하나은행 대출의 원금과 비용. 예상 비용과 실제 납부액을 구분한다."""
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from core import MAX_AMOUNT

HANA_PRINCIPAL = 70_000_000
HANA_SHARES = {'mother': 20_000_000, 'me': 50_000_000}
ANNUAL_RATE = Decimal('0.046')
FEE_RATE = Decimal('0.0059')

def won(value):
    return int(value.quantize(Decimal('1'), rounding=ROUND_HALF_UP))

def monthly_interest(principal):
    return won(Decimal(principal) * ANNUAL_RATE / 12)

def repayment_fee(principal):
    return won(Decimal(principal) * FEE_RATE)

def estimate_fee_rate(record):
    if record['fee_basis']=='actual': return '실제 납부액'
    if record.get('fee_rate'): return record['fee_rate']
    if record['principal']==0: return '율 기록 없음'
    return '0.49%' if record['fee']==won(Decimal(record['principal'])*Decimal('0.0049')) else '0.59%'

def hana_balance(records):
    return HANA_PRINCIPAL - sum(record['principal'] for record in records)

def hana_unallocated(records):
    return [record for record in records if 'mother' not in record]

def hana_balances(records):
    return {person: original-sum(record.get(person,0) for record in records) for person,original in HANA_SHARES.items()}

def validate_hana(record, records, today, allow_legacy=False,allow_historical=False):
    fields = ('id', 'date', 'principal', 'interest', 'fee', 'fee_basis', 'memo')
    if not isinstance(record, dict) or any(field not in record for field in fields):
        raise ValueError('하나은행 상환 내역의 필수 항목이 없습니다.')
    for field in ('id', 'date', 'fee_basis', 'memo'):
        if not isinstance(record[field], str) or ILLEGAL_CHARACTERS_RE.search(record[field]):
            raise ValueError('하나은행 상환 내역의 글이나 특수문자를 확인해주세요.')
    if not record['id'].strip() or len(record['id']) > 100 or len(record['memo']) > 1000:
        raise ValueError('식별 정보나 메모 길이를 확인해주세요.')
    try:
        when = date.fromisoformat(record['date'])
        if when.isoformat() != record['date']:
            raise ValueError()
    except ValueError:
        raise ValueError('상환일 형식을 확인해주세요.') from None
    if when > today:
        raise ValueError('미래 상환은 계산만 할 수 있습니다. 실제 상환 내역에는 저장할 수 없습니다.')
    for field in ('principal', 'interest', 'fee'):
        if type(record[field]) is not int or not 0 <= record[field] <= MAX_AMOUNT:
            raise ValueError('금액은 0 이상의 정수로 입력해주세요.')
    if record['fee_basis'] not in ('estimate', 'actual'):
        raise ValueError('수수료의 예상/실제 구분을 확인해주세요.')
    rate=record.get('fee_rate')
    if rate is not None and rate not in ('0.49%','0.59%'):
        raise ValueError('수수료율 기록을 확인해주세요.')
    if record['fee_basis']=='estimate':
        fees=[repayment_fee(record['principal'])] if rate in (None,'0.59%') else []
        if allow_historical and rate in (None,'0.49%'):
            fees.append(won(Decimal(record['principal'])*Decimal('0.0049')))
        if record['fee'] not in fees:
            raise ValueError('새 예상 수수료는 상환 원금의 0.59%로 계산해야 합니다.')
    if record['principal'] + record['interest'] + record['fee'] <= 0:
        raise ValueError('상환 원금 또는 납부 비용을 입력해주세요.')
    if record['principal'] + record['interest'] + record['fee'] > MAX_AMOUNT:
        raise ValueError('금액이 너무 큽니다.')
    split_fields = ('mother', 'me', 'mother_interest', 'me_interest')
    legacy = all(field not in record for field in split_fields)
    if not (legacy and allow_legacy):
        if any(field not in record or type(record[field]) is not int or not 0 <= record[field] <= MAX_AMOUNT for field in split_fields):
            raise ValueError('엄마·본인 원금과 납부 이자를 각각 입력해주세요.')
        if record['mother'] + record['me'] != record['principal']:
            raise ValueError('엄마 원금 + 본인 원금은 총 상환 원금과 같아야 합니다.')
        if record['mother_interest'] + record['me_interest'] != record['interest']:
            raise ValueError('엄마 이자 + 본인 이자는 실제 납부 이자 합계와 같아야 합니다.')
    others = [r for r in records if r['id'] != record['id']]
    if not legacy:
        remaining = hana_balances(others)
        if any(record[person] > remaining[person] for person in HANA_SHARES):
            raise ValueError('엄마 또는 본인의 하나은행 잔여 원금보다 많이 갚을 수 없습니다.')
        if hana_unallocated(others) and record['id'] not in {r['id'] for r in records}:
            raise ValueError('기존 하나은행 내역의 엄마·본인 배분을 먼저 확인해주세요.')
    if record['principal'] > hana_balance(others):
        raise ValueError('하나은행의 남은 원금보다 많이 갚을 수 없습니다.')

def fold_hana(events):
    records, seen = {}, set()
    for eid, timestamp, kind, raw in events:
        if eid in seen:
            continue
        seen.add(eid)
        if kind == 'hana_save':
            data = json.loads(raw)
            validate_hana(data, [], date.max, allow_legacy=True,allow_historical=True)
            records[data['id']] = data
        elif kind == 'hana_delete':
            data = json.loads(raw)
            if not isinstance(data.get('id'), str) or not data['id']:
                raise ValueError('삭제할 하나은행 내역의 식별 정보가 없습니다.')
            records.pop(data['id'], None)
    result = sorted(records.values(), key=lambda r:(r['date'], r['id']))
    if hana_balance(result) < 0:
        raise ValueError('하나은행 상환 원금 합계가 7천만 원을 초과합니다.')
    if any(value < 0 for value in hana_balances(result).values()):
        raise ValueError('엄마 또는 본인의 하나은행 상환 원금 합계가 각자의 대출액을 초과합니다.')
    return result

def hana_excel(records):
    wb = Workbook()
    summary = wb.active
    summary.title = '하나은행 요약'
    left = hana_balance(records)
    summary.append(['항목', '금액/값'])
    for row in [('대출 원금', HANA_PRINCIPAL), ('상환 원금', HANA_PRINCIPAL-left), ('남은 원금', left), ('연 이자율', '4.6%'), ('수수료율', '0.59%'), ('예상 월 이자', monthly_interest(left)), ('월 이자 감소액', monthly_interest(HANA_PRINCIPAL-left)), ('실제 납부 이자 합계', sum(r['interest'] for r in records)), ('확인된 실제 수수료 합계', sum(r['fee'] for r in records if r['fee_basis']=='actual'))]:
        summary.append(row)
    summary.append(['계산 기준', '월 이자는 잔액 × 연 4.6% ÷ 12. 수수료 예상은 상환액 × 0.59%, 잔여기간 미반영.'])
    shares = wb.create_sheet('엄마·본인 요약')
    shares.append(['구분','빌린 원금','상환 원금','남은 원금','예상 월 이자','실제 납부 이자 합계'])
    balances=hana_balances(records)
    for person,label in [('mother','엄마'),('me','본인')]:
        shares.append([label,HANA_SHARES[person],HANA_SHARES[person]-balances[person],balances[person] if not hana_unallocated(records) else '배분 확인 필요',monthly_interest(balances[person]) if not hana_unallocated(records) else '배분 확인 필요',sum(r.get(person+'_interest',0) for r in records) if not hana_unallocated(records) else '배분 확인 필요'])
    sheet = wb.create_sheet('하나은행 상환 내역')
    sheet.append(['기록 ID','상환일','상환 원금','실제 납부 이자','상환수수료','수수료 구분','메모','엄마 원금','본인 원금','엄마 납부 이자','본인 납부 이자','예상 적용 수수료율'])
    for record in records:
        sheet.append([record['id'],record['date'],record['principal'],record['interest'],record['fee'],'실제' if record['fee_basis']=='actual' else '예상',record['memo']]+[record.get(field,'배분 확인 필요') for field in ('mother','me','mother_interest','me_interest')]+[estimate_fee_rate(record)])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=26
            for cell in column:
                if isinstance(cell.value,str): cell.data_type='s'
                if isinstance(cell.value,int): cell.number_format='#,##0'
    output=BytesIO()
    wb.save(output)
    return output.getvalue()
