from datetime import date, datetime
from zoneinfo import ZoneInfo
from io import BytesIO
import json
import re
from copy import copy
from openpyxl import load_workbook
from core import excel
from hana import hana_excel

def validate_backup(data, today):
    if not isinstance(data,dict) or set(data)!={'date','revision'}:
        raise ValueError('백업 완료 정보를 확인해주세요.')
    if not isinstance(data['date'],str) or not isinstance(data['revision'],str):
        raise ValueError('백업 완료 정보의 형식을 확인해주세요.')
    when=date.fromisoformat(data['date'])
    if when.isoformat()!=data['date'] or when>today or not re.fullmatch(r'[0-9a-f]{64}',data['revision']):
        raise ValueError('백업 완료 일자와 파일 기준 정보를 확인해주세요.')

def next_year(day):
    try:
        return day.replace(year=day.year+1)
    except ValueError:
        return day.replace(year=day.year+1,day=28)

def backup_state(events, today):
    start, last, seen = None, None, set()
    for eid, timestamp, kind, raw in events:
        if eid in seen: continue
        seen.add(eid)
        if kind in ('save','hana_save'):
            try:
                when=datetime.fromisoformat(timestamp).astimezone(ZoneInfo('Asia/Seoul')).date()
            except ValueError:
                when=date.fromisoformat(json.loads(raw)['date'])
            start=min(start,when) if start else when
        elif kind=='backup_confirm':
            data=json.loads(raw)
            validate_backup(data,today)
            when=date.fromisoformat(data['date'])
            last=max(last,when) if last else when
    due=next_year(last or start) if (last or start) else None
    return {'last':last,'start':start,'due':due,'overdue':bool(due and today>=due),'events':events}

def backup_excel(records, plan, hana_records, events, today, revision):
    wb=load_workbook(BytesIO(excel(records,plan,today)))
    bank=load_workbook(BytesIO(hana_excel(hana_records)))
    for source in bank:
        dest=wb.create_sheet(source.title)
        for row in source:
            for cell in row:
                target=dest.cell(cell.row,cell.column,cell.value)
                if cell.has_style: target._style=copy(cell._style)
                if isinstance(cell.value,str): target.data_type='s'
        dest.freeze_panes=source.freeze_panes
        dest.auto_filter.ref=source.auto_filter.ref
        for name,dimension in source.column_dimensions.items():
            dest.column_dimensions[name].width=dimension.width
    meta=wb.create_sheet('백업 정보',0)
    for row in [('항목','내용'),('백업 파일 생성일',today.isoformat()),('가족 유효 기록 수',len(records)),('하나은행 유효 기록 수',len(hana_records)),('원본 이벤트 수',len(events)),('파일 기준 정보',revision),('보관 안내','이 파일을 PC와 다른 저장 위치에 함께 보관하세요. 원본 이벤트 이력도 포함되어 있습니다.')]:
        meta.append(row)
    meta.column_dimensions['A'].width=25
    meta.column_dimensions['B'].width=80
    raw=wb.create_sheet('원본 이벤트 이력')
    raw.append(['event_id','timestamp','type','payload'])
    for row in events: raw.append(row)
    raw.freeze_panes='A2'
    raw.column_dimensions['D'].width=80
    for sheet in (meta,raw):
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str): cell.data_type='s'
    out=BytesIO(); wb.save(out)
    return out.getvalue()
