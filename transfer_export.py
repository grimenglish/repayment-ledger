"""앱이 생성하는 읽기 쉬운 송금 기록. 예상 수수료를 실제 납부액에 합산하지 않는다."""
from datetime import date
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.page import PageMargins
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

HEADERS=['상환일','대출 구분','송금·출금 은행','총 송금·납부액','갚은 원금','실제 납부 이자','실제 수수료','실제 송금자','상환 내용','메모','기록 ID']

def validate_transfer_fields(record):
    for field in ('bank','sender'):
        if field in record and (not isinstance(record[field],str) or len(record[field])>100 or ILLEGAL_CHARACTERS_RE.search(record[field])):
            raise ValueError('송금 은행과 송금자 이름을 확인해주세요.')

def transfer_row(kind,record):
    when=date.fromisoformat(record['date'])
    bank=record.get('bank','').strip() or '은행 미입력'
    sender=record.get('sender','').strip() or '송금자 미입력'
    loan={'save':'가족 대출','hana_save':'하나은행','gf_save':'여자친구 대출'}[kind]
    principal=record['mother']+record['me'] if kind=='save' else record['principal']
    interest=record['interest']
    fee=record['fee'] if kind=='hana_save' and record['fee_basis']=='actual' else '미확인' if kind=='hana_save' else '별도 기록 없음'
    total=record['total'] if kind=='save' else principal+interest+fee if type(fee) is int else '수수료 미확인' if kind=='hana_save' else principal+interest
    words=[f'{when:%Y년 %m월 %d일}',bank,f'{total:,}원 납부' if type(total) is int else f'원금·이자 {principal+interest:,}원 납부 (수수료 미확인)']
    if kind!='gf_save' and 'mother' in record:
        words.append(f'엄마 원금 {record["mother"]:,}원 / 본인 원금 {record["me"]:,}원')
    else:
        words.append(f'원금 {principal:,}원')
    words.append(f'이자 {interest:,}원')
    if kind=='hana_save':
        words.append(f'실제 수수료 {fee:,}원' if type(fee) is int else f'예상 수수료 {record["fee"]:,}원 · 실제 미확인')
    return [when,loan,bank,total,principal,interest,fee,sender,' · '.join(words),record['memo'],record['id']]

def add_transfer_sheet(workbook,title,groups):
    sheet=workbook.create_sheet(title,0)
    sheet.append(HEADERS)
    entries=[(record['date'],record['id'],kind,record) for kind,records in groups for record in records]
    for when,identifier,kind,record in sorted(entries,key=lambda row:row[:3]):
        sheet.append(transfer_row(kind,record))
    sheet.freeze_panes='D2'
    sheet.auto_filter.ref=f'A1:K{max(sheet.max_row,1)}'
    widths=[21,20,24,24,22,22,21,22,65,35,25]
    for i,width in enumerate(widths,1): sheet.column_dimensions[sheet.cell(1,i).column_letter].width=width
    sheet.column_dimensions['K'].hidden=True
    sheet.row_dimensions[1].height=30
    for cell in sheet[1]:
        cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF',size=11)
        cell.fill=PatternFill('solid',fgColor='17233D')
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height=58
        for cell in row:
            cell.font=Font(name='맑은 고딕',size=11,color='17233D')
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            if cell.row%2==0: cell.fill=PatternFill('solid',fgColor='F1F5FA')
            if isinstance(cell.value,str): cell.data_type='s'
            if isinstance(cell.value,int): cell.number_format='#,##0"원"'
        row[0].number_format='yyyy"년" m"월" d"일"'
        for cell in (row[2],row[3],row[7]):
            if isinstance(cell.value,str) and ('미입력' in cell.value or '미확인' in cell.value):
                cell.font=Font(name='맑은 고딕',size=11,color='9A5B00')
    sheet.sheet_view.showGridLines=False
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.page_margins=PageMargins(left=.25,right=.25,top=.4,bottom=.4,header=.2,footer=.2)
    sheet.print_title_rows='1:1';sheet.print_area=f'A1:J{max(sheet.max_row,1)}'
    sheet.oddFooter.center.text='은행 이체확인증·거래내역과 함께 보관하세요.'
    workbook.active=0
    return sheet
