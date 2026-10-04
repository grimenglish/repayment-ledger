import unittest
import json
import time
import threading
from datetime import date,datetime
from zoneinfo import ZoneInfo
from io import BytesIO
from unittest.mock import patch, MagicMock
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from core import *
from storage import GoogleStore, HEADER, service_account_info, clean_rows, KeyFormatError, LedgerDataError, SheetFormatError
from hana import HANA_PRINCIPAL, HANA_SHARES, monthly_interest, repayment_fee, hana_balance, hana_balances, hana_unallocated, hana_excel
from backup import backup_state, backup_excel, next_year
from girlfriend import GF_PRINCIPAL,gf_interest,gf_balance,validate_gf,fold_gf,gf_excel
TODAY=date(2026,10,1)
UI_TODAY=datetime.now(ZoneInfo('Asia/Seoul')).date()
def record(**kw):
    r=dict(id='a',date='2026-10-01',bank='국민',sender='본인',total=5_000_000,mother=4_000_000,me=1_000_000,interest=0,memo='')
    r.update(kw)
    return r
def bank_record(**kw):
    r=dict(id='hana-a',date='2026-10-01',principal=10_000_000,interest=0,fee=59_000,fee_basis='estimate',memo='')
    r.update(kw)
    r.setdefault('mother',0)
    r.setdefault('me',r['principal']-r['mother'])
    r.setdefault('mother_interest',0)
    r.setdefault('me_interest',r['interest']-r['mother_interest'])
    return r
class Sheet:
    def __init__(self): self.rows=[HEADER.copy()]
    def get_all_values(self): return [r.copy() for r in self.rows]
    def append_row(self,row,**kw): self.rows.append(row)
    def col_values(self,col): return [r[col-1] for r in self.rows]
def fake_store():
    s=GoogleStore.__new__(GoogleStore); s.sheet=Sheet(); s.lock=threading.RLock(); return s

def girlfriend_record(**kw):
    r=dict(id='gf-a',date='2026-10-01',principal=5_000_000,interest=50_000,memo='')
    r.update(kw);return r
class Tests(unittest.TestCase):
    def setUp(self):
        import streamlit as st
        st.cache_resource.clear()

    def test_transfer_export_dates_banks_amounts_and_first_sheet(self):
        records=[record(id='b',date='2026-10-04',bank='국민은행',sender='홍길동'),record(id='a',date='2026-10-01',bank='=1+1',memo='=NOW()')]
        wb=load_workbook(BytesIO(excel(records,DEFAULT_PLAN,UI_TODAY)))
        sheet=wb.active
        self.assertEqual(sheet.title,'친척 상환 내역서')
        self.assertEqual(sheet['A2'].value.date(),date(2026,10,1))
        self.assertEqual(sheet['B3'].value,'국민은행')
        self.assertEqual(sheet['C3'].value,5_000_000)
        self.assertEqual(sheet['D3'].value,4_000_000)
        self.assertEqual(sheet['E3'].value,1_000_000)
        self.assertEqual(sheet['G3'].value,190_000_000)
        self.assertEqual(sheet['B2'].data_type,'s')
        self.assertEqual(sheet['I2'].data_type,'s')
        self.assertTrue(sheet.column_dimensions['J'].hidden)
        self.assertEqual(sheet.freeze_panes,'C2')
        self.assertEqual(wb['잔액 요약']['C2'].value,8_000_000)

    def test_family_statement_separate_banks_interest_totals_and_empty(self):
        rows=[record(id='b',bank='신협',date='2026-10-01',mother=0,me=2_000_000,interest=30_000,total=2_030_000),record(id='a',bank='카카오뱅크',date='2026-10-01')]
        wb=load_workbook(BytesIO(excel(rows,DEFAULT_PLAN,TODAY)))
        sheet=wb.active
        self.assertEqual([sheet.cell(r,2).value for r in (2,3)],['카카오뱅크','신협'])
        self.assertEqual([sheet.cell(r,7).value for r in (2,3)],[195_000_000,193_000_000])
        self.assertEqual([sheet.cell(4,c).value for c in range(3,8)],[7_030_000,4_000_000,3_000_000,30_000,193_000_000])
        self.assertEqual(sheet.auto_filter.ref,'A1:I3')
        empty=load_workbook(BytesIO(excel([],DEFAULT_PLAN,TODAY))).active
        self.assertEqual(empty['C2'].value,0)
        self.assertEqual(empty['G2'].value,200_000_000)

    def test_relative_proof_no_person_split_and_exact_transfers(self):
        rows=[record(id='b',bank='신협',mother=0,me=2_000_000,interest=30_000,total=2_030_000),record(id='a',bank='카카오뱅크')]
        wb=load_workbook(BytesIO(relative_proof_excel(rows)))
        self.assertEqual(len(wb.sheetnames),1)
        ws=wb.active
        self.assertEqual([c.value for c in ws[1]],['송금일','보낸 은행','실제 송금액','누적 송금액','상환 후 남은 원금'])
        self.assertEqual([ws.cell(3,c).value for c in range(2,6)],['신협',2_030_000,7_030_000,193_000_000])
        self.assertEqual(ws['C4'].value,7_030_000)
        self.assertEqual(ws['E4'].value,193_000_000)
        self.assertEqual(ws.max_column,5)
        self.assertEqual(ws.auto_filter.ref,'A1:E3')
        empty=load_workbook(BytesIO(relative_proof_excel([]))).active
        self.assertEqual(empty['D2'].value,0)
        self.assertEqual(empty['E2'].value,200_000_000)

    def test_transfer_unknown_bank_and_estimated_fee_are_not_actual(self):
        records=[bank_record(),bank_record(id='b',date='2026-10-04',fee=12_345,fee_basis='actual',bank='신한은행',sender='홍길동')]
        wb=load_workbook(BytesIO(hana_excel(records)));sheet=wb.active
        self.assertEqual(sheet.title,'하나은행 납부 기록')
        self.assertEqual(sheet['C2'].value,'은행 미입력')
        self.assertEqual(sheet['D2'].value,'수수료 미확인')
        self.assertEqual(sheet['G2'].value,'미확인')
        self.assertIn('예상 수수료 59,000원',sheet['I2'].value)
        self.assertEqual(sheet['D3'].value,10_012_345)
        self.assertEqual(sheet['C3'].value,'신한은행')
        gf=load_workbook(BytesIO(gf_excel([girlfriend_record(bank='카카오뱅크',sender='홍길동')])))
        self.assertEqual(gf.active.title,'여자친구 송금 기록')
        self.assertEqual(gf.active['D2'].value,5_050_000)
        self.assertEqual(gf.active['C2'].value,'카카오뱅크')

    def test_transfer_metadata_validation_and_old_records(self):
        s=fake_store()
        old=girlfriend_record()
        s.mutate('gf_save',old,s.load()[2],TODAY)
        self.assertNotIn('bank',s.load_details(include_girlfriend=True)[4][0])
        new=girlfriend_record(bank='우리은행',sender='홍길동')
        s.mutate('gf_save',new,s.load()[2],TODAY)
        self.assertEqual(s.load_details(include_girlfriend=True)[4],[new])
        s.undo_save('gf_save',new,old,s.load()[2],TODAY)
        self.assertEqual(s.load_details(include_girlfriend=True)[4],[old])
        for kind,factory in [('gf_save',girlfriend_record),('hana_save',bank_record)]:
            for changes in ({'bank':1},{'sender':'x\x00'},{'bank':'x'*101}):
                with self.assertRaises(ValueError): s.mutate(kind,factory(**changes),s.load()[2],TODAY)

    def test_combined_transfer_sheet_contains_each_payment(self):
        s=fake_store()
        s.mutate('save',record(bank='국민은행'),s.load()[2],TODAY)
        s.mutate('hana_save',bank_record(bank='하나은행',fee_basis='actual',fee=12_345),s.load()[2],TODAY)
        s.mutate('gf_save',girlfriend_record(bank='카카오뱅크'),s.load()[2],TODAY)
        family,plan,bank,revision,gf,state=s.load_details(include_backup=True,include_girlfriend=True)
        wb=load_workbook(BytesIO(backup_excel(family,plan,bank,state['events'],TODAY,revision)))
        self.assertEqual(wb.active.title,'전체 송금 기록')
        self.assertEqual(wb.active.max_row,4)
        by_id={row[10]:row for row in wb.active.iter_rows(min_row=2,values_only=True)}
        self.assertEqual(by_id['a'][3],5_000_000)
        self.assertEqual(by_id['hana-a'][3],10_012_345)
        self.assertEqual(by_id['gf-a'][2],'카카오뱅크')
        self.assertEqual(len(wb.sheetnames),14)

    def test_ui_transfer_bank_required_and_saved(self):
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.number_input(key='gf_new_principal').set_value(1_000_000).run()
            app.button(key='gf_new_save').click().run()
            self.assertTrue(app.error)
            self.assertEqual(s.load_details(include_girlfriend=True)[4],[])
            app.text_input(key='gf_new_bank').set_value('카카오뱅크')
            app.text_input(key='gf_new_sender').set_value('홍길동')
            app.button(key='gf_new_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load_details(include_girlfriend=True)[4][0]['bank'],'카카오뱅크')
            app.number_input(key='hana_new_me').set_value(1_000_000).run()
            app.button(key='hana_new_save').click().run()
            self.assertEqual(s.load_details()[2],[])
            app.text_input(key='hana_new_bank').set_value('우리은행')
            app.text_input(key='hana_new_sender').set_value('홍길동')
            app.button(key='hana_new_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load_details()[2][0]['bank'],'우리은행')

    def test_girlfriend_interest_and_validation(self):
        self.assertEqual(gf_interest(20_000_000),50_000)
        self.assertEqual(gf_interest(15_000_000),37_500)
        self.assertEqual(gf_interest(18_000_000),45_000)
        self.assertEqual(gf_interest(200),1)
        self.assertEqual(gf_interest(0),0)
        validate_gf(girlfriend_record(),[],TODAY)
        validate_gf(girlfriend_record(principal=0),[],TODAY)
        for invalid in (girlfriend_record(principal=20_000_001),girlfriend_record(principal=True),girlfriend_record(interest=-1),girlfriend_record(principal=0,interest=0),girlfriend_record(date='2026-10-02'),girlfriend_record(date='2026-1-1'),girlfriend_record(memo='x\x00'),girlfriend_record(principal=1.5),girlfriend_record(interest=MAX_AMOUNT)):
            with self.assertRaises(ValueError): validate_gf(invalid,[],TODAY)

    def test_girlfriend_save_edit_delete_and_isolation(self):
        s=fake_store()
        s.mutate('save',record(),s.load()[2],TODAY)
        s.mutate('hana_save',bank_record(),s.load()[2],TODAY)
        before=s.load_details()
        first=girlfriend_record()
        s.mutate('gf_save',first,before[3],TODAY)
        self.assertEqual(s.load_details()[:3],before[:3])
        self.assertEqual(gf_balance(s.load_details(include_girlfriend=True)[4]),15_000_000)
        s.mutate('gf_save',girlfriend_record(principal=10_000_000),s.load()[2],TODAY)
        self.assertEqual(gf_balance(s.load_details(include_girlfriend=True)[4]),10_000_000)
        with self.assertRaises(ValueError): s.mutate('gf_save',girlfriend_record(id='gf-b',principal=10_000_001),s.load()[2],TODAY)
        s.mutate('gf_delete',{'id':first['id']},s.load()[2],TODAY)
        self.assertEqual(s.load_details(include_girlfriend=True)[4],[])
        self.assertEqual(s.load_details()[:3],before[:3])

    def test_girlfriend_undo_and_conflict(self):
        s=fake_store();saved=girlfriend_record()
        s.mutate('gf_save',saved,s.load()[2],TODAY)
        old_revision=s.load()[2]
        edited=girlfriend_record(principal=10_000_000)
        s.mutate('gf_save',edited,old_revision,TODAY)
        with self.assertRaises(ValueError): s.undo_save('gf_save',saved,None,old_revision,TODAY)
        s.undo_save('gf_save',edited,saved,s.load()[2],TODAY)
        self.assertEqual(s.load_details(include_girlfriend=True)[4],[saved])
        s.undo_save('gf_save',saved,None,s.load()[2],TODAY)
        self.assertEqual(s.load_details(include_girlfriend=True)[4],[])

    def test_girlfriend_backup_and_event_replay(self):
        s=fake_store();saved=girlfriend_record(memo='=1+1')
        s.mutate('gf_save',saved,s.load()[2],TODAY)
        records,plan,bank,revision,gf,state=s.load_details(include_backup=True,include_girlfriend=True)
        wb=load_workbook(BytesIO(backup_excel(records,plan,bank,state['events'],TODAY,revision)))
        self.assertEqual(wb['여자친구 상환 내역']['C2'].value,5_000_000)
        self.assertEqual(wb['여자친구 상환 내역']['H2'].data_type,'s')
        events=[list(row) for row in wb['원본 이벤트 이력'].iter_rows(min_row=2,values_only=True)]
        self.assertEqual(fold_gf(events),gf)
        timestamp=state['events'][0][1]
        from datetime import datetime
        began=datetime.fromisoformat(timestamp).astimezone(ZoneInfo('Asia/Seoul')).date()
        self.assertEqual(state['start'],began)
        wb=load_workbook(BytesIO(gf_excel(gf)))
        self.assertEqual(wb['여자친구 대출 요약']['B6'].value,37_500)

    def test_girlfriend_corrupt_cumulative_records(self):
        s=fake_store()
        for i in range(2):
            s.sheet.rows.append([str(i),'2026-10-01T00:00:00+00:00','gf_save',json.dumps(girlfriend_record(id=str(i),principal=15_000_000))])
        with self.assertRaises(LedgerDataError): s.load_details()

    def test_hana_historical_fees_preserved_and_new_rate_enforced(self):
        from hana import estimate_fee_rate
        s=fake_store();old=bank_record(fee=49_000)
        s.sheet.rows.append(['old','2026-10-01T00:00:00+00:00','hana_save',json.dumps(old)])
        self.assertEqual(s.load_details()[2],[old])
        self.assertEqual(estimate_fee_rate(old),'0.49%')
        with self.assertRaises(ValueError): s.mutate('hana_save',bank_record(id='new',fee=49_000),s.load()[2],TODAY)
        updated=bank_record(fee_rate='0.59%')
        s.mutate('hana_save',updated,s.load()[2],TODAY)
        self.assertEqual(s.load_details()[2],[updated])
        s.undo_save('hana_save',updated,old,s.load()[2],TODAY)
        self.assertEqual(s.load_details()[2],[old])
        wb=load_workbook(BytesIO(hana_excel([old])))
        self.assertEqual(wb['하나은행 상환 내역']['E2'].value,49_000)
        self.assertEqual(wb['하나은행 상환 내역']['L2'].value,'0.49%')
        actual=bank_record(fee=12_345,fee_basis='actual')
        s.mutate('hana_save',actual,s.load()[2],TODAY)
        self.assertEqual(s.load_details()[2][0]['fee'],12_345)

    def test_girlfriend_ui_save_edit_delete_and_undo(self):
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            self.assertFalse(app.exception)
            self.assertTrue(any(m.label=='내 남은 대출' and m.value=='170,000,000원' for m in app.metric))
            app.button(key='card_me_gf').click().run()
            self.assertEqual(app.session_state['main_navigation'],'여자친구 대출')
            self.assertEqual(app.session_state['gf_navigation'],'여자친구 상환 입력')
            app.text_input(key='gf_new_bank').set_value('국민은행')
            app.button(key='gf_new_save').click().run()
            self.assertTrue(app.error)
            app.number_input(key='gf_new_principal').set_value(5_000_000).run()
            app.number_input(key='gf_new_interest').set_value(50_000).run()
            app.text_input(key='gf_new_bank').set_value('국민은행')
            app.button(key='gf_new_save').click().run()
            self.assertFalse(app.exception)
            gf=s.load_details(include_girlfriend=True)[4]
            self.assertEqual(len(gf),1)
            self.assertEqual(gf_balance(gf),15_000_000)
            self.assertTrue(any('약 12,500원 감소' in m.value for m in app.markdown))
            self.assertTrue(any(m.label=='내 남은 대출' and m.value=='165,000,000원' for m in app.metric))
            identifier=gf[0]['id']
            app.button(key='gf_history_edit_'+identifier).click().run()
            app.number_input(key='gf_edit_'+identifier+'_principal').set_value(10_000_000).run()
            app.text_input(key='gf_edit_'+identifier+'_bank').set_value('국민은행')
            app.button(key='gf_edit_'+identifier+'_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(gf_balance(s.load_details(include_girlfriend=True)[4]),10_000_000)
            app.button(key='undo_last_save').click().run()
            self.assertEqual(gf_balance(s.load_details(include_girlfriend=True)[4]),15_000_000)
            app.checkbox(key='gf_confirm_'+identifier).set_value(True).run()
            app.button(key='gf_delete_'+identifier).click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load_details(include_girlfriend=True)[4],[])

    def test_girlfriend_simulation_full_payment_and_reload(self):
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.number_input(key='gf_simulation').set_value(20_000_000).run()
            self.assertEqual(s.load_details(include_girlfriend=True)[4],[])
            app.button(key='gf_new_all').click().run()
            app.text_input(key='gf_new_bank').set_value('국민은행')
            app.button(key='gf_new_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(gf_balance(s.load_details(include_girlfriend=True)[4]),0)
            self.assertTrue(any('여자친구 대출 완납' in m.value for m in app.markdown))
            reloaded=AppTest.from_file('app.py')
            reloaded.secrets['login']={'salt':'test','password_hash':'test'}
            reloaded.session_state['authenticated_until']=time.time()+1000
            reloaded.run()
            self.assertFalse(reloaded.exception)
            self.assertTrue(reloaded.button(key='card_me_gf').disabled)

    def test_monthly_goal_states(self):
        from design_ui import monthly_goal
        self.assertEqual(monthly_goal(96_000_000,4_000_000,4_000_000)['status'],'이번 달 목표 달성')
        self.assertEqual(monthly_goal(99_000_000,1_000_000,4_000_000)['due'],3_000_000)
        self.assertEqual(monthly_goal(500_000,0,4_000_000)['due'],500_000)
        self.assertEqual(monthly_goal(5_000_000,0,0)['status'],'월 목표 설정 필요')
        self.assertEqual(monthly_goal(0,100_000_000,4_000_000)['status'],'완납 완료')

    def test_milestone_only_new_verified_repayments(self):
        from design_ui import make_receipt, achievement_messages
        saved=record(mother=10_000_000,me=0,total=10_000_000)
        value=make_receipt('save',saved,None,[],[])
        self.assertIn('10% 상환 달성',achievement_messages(value)[0])
        self.assertEqual(achievement_messages(make_receipt('save',saved,saved,[saved],[])),[])
        self.assertEqual(achievement_messages(make_receipt('save',record(mother=0,me=0,total=1,interest=1),None,[],[])),[])
        bank=bank_record(principal=20_000_000,mother=20_000_000,me=0)
        self.assertIn('완납',achievement_messages(make_receipt('hana_save',bank,None,[],[]))[0])

    def test_person_card_routes_both_loans_and_does_not_save(self):
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.toggle(key='family_auto_fill').set_value(True).run()
            app.number_input(key='new_quick_me').set_value(500_000).run()
            app.button(key='card_mother_family').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state['main_navigation'],'가족 대출')
            self.assertEqual(app.radio(key='new_manual_person').value,'엄마만')
            self.assertEqual(app.number_input(key='new_manual_mother').value,0)
            self.assertFalse(app.toggle(key='family_auto_fill').value)
            app.button(key='card_me_hana').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state['main_navigation'],'하나은행 대출')
            self.assertEqual(app.session_state['hana_navigation'],'하나은행 상환 입력')
            self.assertEqual(app.radio(key='hana_new_person').value,'본인만')
            self.assertEqual(app.number_input(key='hana_new_me').value,0)
            self.assertEqual(s.load_details()[0],[])
            self.assertEqual(s.load_details()[2],[])

    def test_dashboard_month_goal_visible(self):
        s=fake_store()
        s.mutate('save',record(),s.load()[2],TODAY)
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            self.assertFalse(app.exception)
            panels=[m.value for m in app.markdown if 'class="monthly-goal"' in m.value]
            self.assertEqual(len(panels),2)
            self.assertTrue(all('이번 달 목표 달성' in m for m in panels))

    def test_quote_pool_and_no_immediate_repeat(self):
        from quotes_ui import QUOTES, choose_quote
        self.assertEqual(len(QUOTES),7)
        for index in range(len(QUOTES)):
            self.assertNotEqual(choose_quote(index),index)
            self.assertTrue(all(QUOTES[index]))
            self.assertTrue(QUOTES[index][-1].startswith('https://www.gutenberg.org/'))

    def test_header_quote_stays_during_input_and_rotates(self):
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            before=app.session_state['header_quote']
            app.number_input(key='new_manual_me').set_value(100_000).run()
            self.assertEqual(app.session_state['header_quote'],before)
            app.button(key='rotate_header_quote').click().run()
            self.assertFalse(app.exception)
            self.assertNotEqual(app.session_state['header_quote'],before)
            self.assertEqual(app.number_input(key='new_manual_me').value,100_000)
            self.assertEqual(s.load()[0],[])

    def test_receipt_balances_new_and_edit(self):
        from design_ui import make_receipt
        old=record()
        value=make_receipt('save',old,None,[],[])
        self.assertEqual(value['before']['mother']-value['after']['mother'],4_000_000)
        changed=record(mother=3_000_000,total=4_000_000)
        value=make_receipt('save',changed,old,[old],[])
        self.assertTrue(value['edited'])
        self.assertEqual(value['after']['mother']-value['before']['mother'],1_000_000)

    def test_bank_receipt_interest_and_legacy(self):
        from design_ui import make_receipt
        saved=bank_record()
        value=make_receipt('hana_save',saved,None,[],[])
        self.assertEqual(monthly_interest(value['before']['me'])-monthly_interest(value['after']['me']),38_334)
        legacy={key:item for key,item in saved.items() if key not in ('mother','me','mother_interest','me_interest')}
        value=make_receipt('hana_save',saved,legacy,[],[legacy])
        self.assertFalse(value['before_known'])
        self.assertTrue(value['split_ready'])

    def test_design_navigation_and_history_edit(self):
        s=fake_store()
        s.mutate('save',record(),s.load()[2],TODAY)
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.radio(key='start_repayment_loan').set_value('하나은행 대출').run()
            app.button(key='start_repayment').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state['main_navigation'],'하나은행 대출')
            app.button(key='family_history_edit_a').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.selectbox(key='family_edit_selection').value,'a')
            bars=[m.value for m in app.markdown if 'role="progressbar"' in m.value]
            self.assertTrue(any('aria-valuenow="3.3"' in m and '#0d9488' in m for m in bars))

    def test_design_saved_result(self):
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.secrets['sender_name']='본인'
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.radio(key='new_manual_person').set_value('본인만').run()
            app.number_input(key='new_manual_me').set_value(1_000_000).run()
            app.button(key='new_manual_save').click().run()
            self.assertFalse(app.exception)
            results=[m.value for m in app.markdown if 'class="payment-receipt"' in m.value]
            self.assertEqual(len(results),1)
            self.assertIn('본인 1,000,000원 상환 완료',results[0])
            self.assertIn('100,000,000원 → 99,000,000원',results[0])

    def test_annual_backup_dates(self):
        value=['e','2026-10-01T01:00:00+00:00','save',json.dumps(record())]
        self.assertFalse(backup_state([value],date(2027,9,30))['overdue'])
        self.assertTrue(backup_state([value],date(2027,10,1))['overdue'])
        done=['b','2027-10-01T01:00:00+00:00','backup_confirm',json.dumps({'date':'2027-10-01','revision':'a'*64})]
        status=backup_state([value,done],date(2027,10,1))
        self.assertFalse(status['overdue'])
        self.assertEqual(status['due'],date(2028,10,1))
        self.assertEqual(next_year(date(2028,2,29)),date(2029,2,28))
        self.assertIsNone(backup_state([],TODAY)['due'])

    def test_full_backup_includes_replayable_events(self):
        from hana import fold_hana
        s=fake_store()
        s.mutate('save',record(memo='=1+1'),s.load()[2],TODAY)
        s.mutate('hana_save',bank_record(),s.load()[2],TODAY)
        records,plan,bank,revision,state=s.load_details(include_backup=True)
        wb=load_workbook(BytesIO(backup_excel(records,plan,bank,state['events'],TODAY,revision)))
        self.assertEqual(len(wb.sheetnames),14)
        self.assertEqual(wb['상환 내역']['I2'].data_type,'s')
        self.assertEqual(wb['하나은행 상환 내역']['C2'].value,10_000_000)
        events=[list(row) for row in wb['원본 이벤트 이력'].iter_rows(min_row=2,values_only=True)]
        self.assertEqual(fold(events),(records,plan))
        self.assertEqual(fold_hana(events),bank)

    def test_backup_confirmation_keeps_records(self):
        s=fake_store()
        s.mutate('save',record(),s.load()[2],TODAY)
        old=s.load_details(); revision=old[3]
        with self.assertRaises(ValueError): s.mutate('backup_confirm',{'date':TODAY.isoformat(),'revision':'a'*64},revision,TODAY)
        s.mutate('backup_confirm',{'date':TODAY.isoformat(),'revision':revision},revision,TODAY)
        current=s.load_details(include_backup=True)
        self.assertEqual(current[:3],old[:3])
        self.assertEqual(current[4]['last'],TODAY)
        self.assertEqual(current[4]['due'],date(2027,10,1))

    def test_backup_ui_confirm_requires_current_snapshot(self):
        import streamlit as st
        st.cache_resource.clear(); s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            self.assertNotIn('confirm_full_backup',[button.key for button in app.button])
            app.session_state['backup_downloaded']={'revision':'a'*64,'date':UI_TODAY.isoformat()}
            app.run()
            self.assertNotIn('confirm_full_backup',[button.key for button in app.button])
            app.session_state['backup_downloaded']={'revision':s.load()[2],'date':UI_TODAY.isoformat()}
            app.run()
            app.button(key='confirm_full_backup').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load_details(include_backup=True)[4]['last'],UI_TODAY)
            self.assertEqual(s.load_details()[0],[])

    def test_undo_new_and_edited_records(self):
        for kind,first,edited in [('save',record(),record(me=2_000_000,total=6_000_000)),('hana_save',bank_record(),bank_record(principal=5_000_000,fee=29_500))]:
            s=fake_store()
            s.mutate(kind,first,s.load()[2],TODAY)
            s.undo_save(kind,first,None,s.load()[2],TODAY)
            self.assertEqual(s.load_details()[0 if kind=='save' else 2],[])
            s.mutate(kind,first,s.load()[2],TODAY)
            s.mutate(kind,edited,s.load()[2],TODAY)
            s.undo_save(kind,edited,first,s.load()[2],TODAY)
            self.assertEqual(s.load_details()[0 if kind=='save' else 2],[first])

    def test_undo_blocks_later_changes(self):
        s=fake_store(); first=record(); second=record(me=2_000_000,total=6_000_000)
        s.mutate('save',first,s.load()[2],TODAY)
        revision=s.load()[2]
        s.mutate('save',second,revision,TODAY)
        with self.assertRaises(ValueError): s.undo_save('save',first,None,revision,TODAY)
        with self.assertRaises(ValueError): s.undo_save('save',first,None,s.load()[2],TODAY)
        self.assertEqual(s.load()[0],[second])

    def test_undo_restores_legacy_bank_record(self):
        s=fake_store(); legacy=bank_record()
        for key in ('mother','me','mother_interest','me_interest'): legacy.pop(key)
        s.sheet.rows.append(['old','time','hana_save',json.dumps(legacy)])
        allocated=bank_record(mother=2_000_000)
        s.mutate('hana_save',allocated,s.load()[2],TODAY)
        s.undo_save('hana_save',allocated,legacy,s.load()[2],TODAY)
        self.assertEqual(s.load_details()[2],[legacy])

    def test_easy_family_repeat_person_buttons_and_undo(self):
        import streamlit as st
        st.cache_resource.clear(); s=fake_store()
        previous=record(mother=0,me=1_000_000,total=1_000_000,date='2026-09-30')
        s.mutate('save',previous,s.load()[2],TODAY)
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            self.assertTrue(any(m.label=='내 남은 대출' and m.value=='169,000,000원' for m in app.metric))
            app.button(key='new_manual_repeat').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.radio(key='new_manual_person').value,'본인만')
            self.assertNotIn('new_manual_mother',[widget.key for widget in app.number_input])
            self.assertEqual(app.number_input(key='new_manual_me').value,1_000_000)
            self.assertEqual(app.date_input(key='new_manual_date').value,UI_TODAY)
            app.button(key='new_manual_me_plus100k').click().run()
            self.assertEqual(app.number_input(key='new_manual_me').value,1_100_000)
            app.button(key='new_manual_me_plus1m').click().run()
            self.assertEqual(app.number_input(key='new_manual_me').value,2_100_000)
            app.button(key='new_manual_me_reset').click().run()
            self.assertEqual(app.number_input(key='new_manual_me').value,0)
            app.button(key='new_manual_repeat').click().run()
            app.button(key='new_manual_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(s.load()[0]),2)
            app.button(key='undo_last_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load()[0],[previous])
            self.assertNotIn('undo_last_save',[button.key for button in app.button])

    def test_easy_hana_repeat_and_undo(self):
        import streamlit as st
        st.cache_resource.clear(); s=fake_store()
        previous=bank_record(date='2026-09-30',mother=10_000_000,me=0)
        s.mutate('hana_save',previous,s.load()[2],TODAY)
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.button(key='hana_new_repeat').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.radio(key='hana_new_person').value,'엄마만')
            self.assertNotIn('hana_new_me',[widget.key for widget in app.number_input])
            self.assertEqual(app.number_input(key='hana_new_mother').value,10_000_000)
            app.button(key='hana_new_mother_plus1m').click().run()
            self.assertEqual(app.number_input(key='hana_new_mother').value,10_000_000)
            app.text_input(key='hana_new_bank').set_value('국민은행')
            app.button(key='hana_new_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(hana_balances(s.load_details()[2])['mother'],0)
            app.button(key='undo_last_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load_details()[2],[previous])

    def test_repeat_limits_remaining_principal(self):
        import streamlit as st
        st.cache_resource.clear(); s=fake_store()
        s.mutate('hana_save',bank_record(mother=15_000_000,me=0,principal=15_000_000,fee=88_500),s.load()[2],TODAY)
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.button(key='hana_new_repeat').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.number_input(key='hana_new_mother').value,5_000_000)
            self.assertTrue(any('남은 원금에 맞춰' in info.value for info in app.info))

    def test_hana_mother_and_me_separate(self):
        s=fake_store()
        self.assertEqual(hana_balances([]),HANA_SHARES)
        self.assertEqual(monthly_interest(20_000_000),76_667)
        self.assertEqual(monthly_interest(50_000_000),191_667)
        s.mutate('hana_save',bank_record(principal=5_000_000,mother=5_000_000,fee=29_500),s.load()[2],TODAY)
        self.assertEqual(hana_balances(s.load_details()[2]),{'mother':15_000_000,'me':50_000_000})
        s.mutate('hana_save',bank_record(id='b',principal=10_000_000,mother=0,fee=59_000),s.load()[2],TODAY)
        self.assertEqual(hana_balances(s.load_details()[2]),{'mother':15_000_000,'me':40_000_000})
        self.assertEqual(monthly_interest(hana_balances(s.load_details()[2])['mother']),57_500)
        self.assertEqual(monthly_interest(hana_balances(s.load_details()[2])['me']),153_333)

    def test_hana_split_validation(self):
        s=fake_store()
        for value in (bank_record(mother=11_000_000,me=0),bank_record(principal=21_000_000,mother=21_000_000,fee=123_900),bank_record(principal=51_000_000,fee=300_900),bank_record(mother_interest=1,me_interest=0),bank_record(mother=True)):
            with self.assertRaises(ValueError): s.mutate('hana_save',value,s.load()[2],TODAY)

    def test_hana_legacy_preserved_and_allocated(self):
        s=fake_store(); legacy=bank_record()
        for key in ('mother','me','mother_interest','me_interest'): legacy.pop(key)
        s.sheet.rows.append(['old','time','hana_save',json.dumps(legacy)])
        bank=s.load_details()[2]
        self.assertEqual(hana_balance(bank),60_000_000)
        self.assertEqual(len(hana_unallocated(bank)),1)
        with self.assertRaises(ValueError): s.mutate('hana_save',bank_record(id='new'),s.load()[2],TODAY)
        s.mutate('hana_save',bank_record(mother=2_000_000),s.load()[2],TODAY)
        self.assertEqual(hana_unallocated(s.load_details()[2]),[])
        self.assertEqual(hana_balances(s.load_details()[2]),{'mother':18_000_000,'me':42_000_000})

    def test_hana_legacy_ui_allocation(self):
        import streamlit as st
        st.cache_resource.clear(); s=fake_store(); legacy=bank_record(interest=200_000)
        for key in ('mother','me','mother_interest','me_interest'): legacy.pop(key)
        s.sheet.rows.append(['old','time','hana_save',json.dumps(legacy)])
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            self.assertFalse(app.exception)
            app.number_input(key='hana_edit_hana-a_mother').set_value(2_000_000)
            app.number_input(key='hana_edit_hana-a_me').set_value(8_000_000)
            app.number_input(key='hana_edit_hana-a_mother_interest').set_value(50_000)
            app.number_input(key='hana_edit_hana-a_me_interest').set_value(150_000)
            app.text_input(key='hana_edit_hana-a_bank').set_value('국민은행')
            app.button(key='hana_edit_hana-a_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(s.load_details()[2]),1)
            self.assertEqual(hana_unallocated(s.load_details()[2]),[])
            self.assertEqual(hana_balances(s.load_details()[2]),{'mother':18_000_000,'me':42_000_000})

    def test_hana_cost_calculation(self):
        self.assertEqual(monthly_interest(70_000_000),268_333)
        self.assertEqual(monthly_interest(60_000_000),230_000)
        self.assertEqual(monthly_interest(10_000_000),38_333)
        self.assertEqual(repayment_fee(10_000_000),59_000)
        self.assertEqual(repayment_fee(20_000_000),118_000)
        self.assertEqual(repayment_fee(70_000_000),413_000)
        self.assertEqual(monthly_interest(0),0)

    def test_hana_save_edit_delete_isolation(self):
        s=fake_store(); s.mutate('save',record(),s.load()[2],TODAY)
        family,plan,_,_=s.load_details()
        s.mutate('hana_save',bank_record(),s.load()[2],TODAY)
        current_family,current_plan,bank,revision=s.load_details()
        self.assertEqual((family,plan),(current_family,current_plan))
        self.assertEqual(hana_balance(bank),60_000_000)
        s.mutate('hana_save',bank_record(principal=20_000_000,fee=118_000),revision,TODAY)
        self.assertEqual(hana_balance(s.load_details()[2]),50_000_000)
        s.mutate('hana_delete',{'id':'hana-a'},s.load()[2],TODAY)
        self.assertEqual(s.load_details()[2],[])
        self.assertEqual(s.load()[0],family)

    def test_hana_invalid_and_overpayment(self):
        s=fake_store()
        for invalid in (bank_record(principal=80_000_000,fee=472_000),bank_record(fee=-1),bank_record(fee=1),bank_record(principal=True),bank_record(date='2026-10-02'),bank_record(memo='bad\x00')):
            with self.assertRaises(ValueError): s.mutate('hana_save',invalid,s.load()[2],TODAY)
        s.mutate('hana_save',bank_record(principal=60_000_000,mother=20_000_000,fee=354_000),s.load()[2],TODAY)
        with self.assertRaises(ValueError): s.mutate('hana_save',bank_record(id='b',principal=20_000_000,fee=118_000),s.load()[2],TODAY)
        self.assertEqual(hana_balance(s.load_details()[2]),10_000_000)

    def test_hana_actual_costs_and_excel(self):
        s=fake_store()
        s.mutate('hana_save',bank_record(interest=200_000,fee=12_345,fee_basis='actual',memo='=1+1'),s.load()[2],TODAY)
        records=s.load_details()[2]
        self.assertEqual(hana_balance(records),60_000_000)
        wb=load_workbook(BytesIO(hana_excel(records)))
        self.assertEqual(wb['하나은행 상환 내역']['E2'].value,12_345)
        self.assertEqual(wb['하나은행 상환 내역']['G2'].data_type,'s')
        self.assertEqual(wb['하나은행 요약']['B4'].value,60_000_000)

    def test_hana_ui_save_and_simulation(self):
        import streamlit as st
        st.cache_resource.clear(); s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.number_input(key='hana_simulation_me').set_value(20_000_000).run()
            self.assertEqual(s.load_details()[2],[])
            app.number_input(key='hana_new_me').set_value(10_000_000).run()
            self.assertTrue(any(m.value=='59,000원' for m in app.metric))
            app.text_input(key='hana_new_bank').set_value('국민은행')
            app.button(key='hana_new_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(hana_balance(s.load_details()[2]),60_000_000)
            self.assertEqual(s.load()[0],[])
            self.assertTrue(any(m.value=='약 230,000원' for m in app.metric))
            record_id=s.load_details()[2][0]['id']
            app.number_input(key='hana_edit_'+record_id+'_me').set_value(5_000_000).run()
            app.text_input(key='hana_edit_'+record_id+'_bank').set_value('국민은행')
            app.button(key='hana_edit_'+record_id+'_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(hana_balance(s.load_details()[2]),65_000_000)
            app.checkbox(key='hana_confirm_'+record_id).set_value(True).run()
            app.button(key='hana_delete_'+record_id).click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load_details()[2],[])

    def test_hana_full_repayment_ui(self):
        import streamlit as st
        st.cache_resource.clear(); s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            app.number_input(key='hana_new_mother').set_value(20_000_000)
            app.number_input(key='hana_new_me').set_value(50_000_000).run()
            app.text_input(key='hana_new_bank').set_value('국민은행')
            app.button(key='hana_new_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(hana_balance(s.load_details()[2]),0)
            self.assertTrue(any(m.value=='약 0원' for m in app.metric))

    def test_korean_amounts(self):
        for amount, expected in [(0,'영 원'), (100_000,'십만 원'), (1_000_000,'백만 원'), (10_000_000,'천만 원'), (4_000_000,'사백만 원'), (100_000_000,'일억 원'), (123_456_789,'일억 이천삼백사십오만 육천칠백팔십구 원')]:
            self.assertEqual(amount_words(amount),expected)

    def test_live_amounts_and_transfer_comparison(self):
        import streamlit as st
        st.cache_resource.clear()
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            for widget in app.text_input:
                if widget.label=='송금 은행': widget.set_value('신한은행')
                if widget.label=='실제 송금자명': widget.set_value('본인')
            for widget in app.number_input:
                if widget.label=='엄마 원금 상환액': widget.set_value(4_000_000)
                if widget.label=='본인 원금 상환액': widget.set_value(100_000)
            app.run()
            self.assertFalse(app.exception)
            self.assertTrue(any(c.value=='100,000원 · 십만 원' for c in app.caption))
            self.assertTrue(any(m.value=='4,100,000원' for m in app.metric))
            app.toggle(key='new_manual_compare').set_value(True).run()
            app.number_input(key='new_manual_total').set_value(5_000_000).run()
            app.button(key='new_manual_save').click().run()
            self.assertTrue(app.error)
            self.assertEqual(s.load()[0],[])
            app.toggle(key='new_manual_compare').set_value(False).run()
            app.number_input(key='new_manual_me').set_value(1_000_000).run()
            self.assertTrue(any(c.value=='1,000,000원 · 백만 원' for c in app.caption))
            self.assertTrue(any(m.value=='5,000,000원' for m in app.metric))
            app.button(key='new_manual_save').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(s.load()[0][0]['total'],5_000_000)

    def test_google_login_blocks_other_accounts(self):
        import streamlit as st
        class User(dict):
            is_logged_in=True
        for email,verified in [('other@example.com',True),('me@example.com',False)]:
            st.cache_resource.clear()
            with patch('streamlit.user',User(email=email,email_verified=verified)), patch('storage.GoogleStore') as connection:
                app=AppTest.from_file('app.py')
                app.secrets['login_mode']='google'
                app.secrets['allowed_email']='me@example.com'
                app.secrets['auth']={'google':{}}
                app.run()
                self.assertFalse(app.exception)
                self.assertTrue(app.error)
                connection.assert_not_called()

    def test_google_login_allows_verified_owner(self):
        import streamlit as st
        class User(dict):
            is_logged_in=True
        st.cache_resource.clear()
        with patch('streamlit.user',User(email='ME@example.com',email_verified=True)), patch('storage.GoogleStore',return_value=fake_store()) as connection:
            app=AppTest.from_file('app.py')
            app.secrets['login_mode']='google'
            app.secrets['allowed_email']='me@example.com'
            app.secrets['auth']={'google':{}}
            app.run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.tabs),12)
            connection.assert_called_once()

    def test_paid_dashboard_with_zero_plan(self):
        import streamlit as st
        st.cache_resource.clear()
        s=fake_store()
        s.mutate('save',record(mother=100_000_000,me=100_000_000,total=200_000_000),s.load()[2],TODAY)
        s.mutate('plan',{'mother':0,'me':0},s.load()[2],TODAY)
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.get('plotly_chart')),6)
            self.assertEqual(len(app.dataframe),1)

    def test_invalid_saved_data(self):
        cases = [[], {'id':'a'}, record(me='100'), record(total=True), record(date='20261001'), record(date='wrong'), record(bank=None), record(memo='bad\x00text'), record(id='')]
        for value in cases:
            with self.subTest(value=value):
                s = fake_store()
                s.sheet.rows.append(['e','time','save',json.dumps(value)])
                with self.assertRaises(LedgerDataError): s.load()

    def test_invalid_saved_plan(self):
        for plan in ({}, {'mother':-1,'me':1}, {'mother':True,'me':1}, {'mother':1,'me':'2'}, {'mother':MAX_AMOUNT+1,'me':0}):
            s=fake_store()
            s.sheet.rows.append(['e','time','plan',json.dumps(plan)])
            with self.assertRaises(LedgerDataError): s.load()

    def test_corrupt_json_and_events(self):
        for row in (['e','time','save','broken'], ['e','time','unknown','{}'], ['e','time'], ['', 'time','plan',json.dumps(DEFAULT_PLAN)]):
            s=fake_store(); s.sheet.rows.append(row)
            with self.assertRaises(LedgerDataError): s.load()

    def test_overpayment_from_saved_events(self):
        s=fake_store()
        for i in range(2):
            value=record(id=str(i),mother=60_000_000,me=0,total=60_000_000)
            s.sheet.rows.append([str(i),'time','save',json.dumps(value)])
        with self.assertRaises(LedgerDataError): s.load()

    def test_edit_cannot_exceed_remaining(self):
        s=fake_store()
        s.mutate('save',record(id='a',mother=60_000_000,me=0,total=60_000_000),s.load()[2],TODAY)
        s.mutate('save',record(id='b',mother=30_000_000,me=0,total=30_000_000),s.load()[2],TODAY)
        with self.assertRaises(ValueError):
            s.mutate('save',record(id='a',mother=80_000_000,me=0,total=80_000_000),s.load()[2],TODAY)
        self.assertEqual(balances(s.load()[0])['mother'],10_000_000)

    def test_response_loss_after_success(self):
        s=fake_store(); original=s.sheet.append_row
        def lost(row, **kwargs):
            original(row,**kwargs)
            raise TimeoutError('response lost')
        s.sheet.append_row=lost
        s.mutate('save',record(),s.load()[2],TODAY)
        self.assertEqual(len(s.sheet.rows),2)
        self.assertEqual(len(s.load()[0]),1)

    def test_failure_before_save(self):
        s=fake_store()
        s.sheet.append_row=MagicMock(side_effect=TimeoutError())
        with self.assertRaises(TimeoutError): s.mutate('save',record(),s.load()[2],TODAY)
        self.assertEqual(s.load()[0],[])

    def test_concurrent_saves(self):
        s=fake_store(); revision=s.load()[2]; results=[]
        def save(i):
            try: s.mutate('save',record(id=str(i)),revision,TODAY); results.append('saved')
            except ValueError: results.append('conflict')
        workers=[threading.Thread(target=save,args=(i,)) for i in range(2)]
        for worker in workers: worker.start()
        for worker in workers: worker.join()
        self.assertCountEqual(results,['saved','conflict'])
        self.assertEqual(len(s.load()[0]),1)

    def test_duplicate_event_replay(self):
        value=['e','time','save',json.dumps(record())]
        rows,plan=fold([value,value])
        self.assertEqual(len(rows),1)

    def test_connection_initialization(self):
        import gspread
        for existing in ([], [HEADER + ['', '']], [['wrong']]):
            client=MagicMock(); sheet=client.open_by_key.return_value.worksheet.return_value
            sheet.get_all_values.return_value=existing
            with patch('storage.service_account_info',return_value={}), patch('storage.Credentials.from_service_account_info'), patch('storage.gspread.authorize',return_value=client):
                config={'gcp_service_account':{}, 'spreadsheet_id':' https://docs.google.com/spreadsheets/d/test-id/edit '}
                if existing == [['wrong']]:
                    with self.assertRaises(SheetFormatError): GoogleStore(config)
                    sheet.append_row.assert_not_called()
                else:
                    GoogleStore(config)
                    client.open_by_key.assert_called_once_with('test-id')
                    client.set_timeout.assert_called_once_with((10,30))
                    if not existing: sheet.append_row.assert_called_once_with(HEADER,value_input_option='RAW')
        client=MagicMock(); book=client.open_by_key.return_value
        book.worksheet.side_effect=gspread.WorksheetNotFound()
        book.add_worksheet.return_value.get_all_values.return_value=[]
        with patch('storage.service_account_info',return_value={}), patch('storage.Credentials.from_service_account_info'), patch('storage.gspread.authorize',return_value=client):
            GoogleStore({'gcp_service_account':{},'spreadsheet_id':'id'})
        book.add_worksheet.assert_called_once_with('ledger_events_v1',rows=2000,cols=4)

    def test_safe_connection_error_screen(self):
        import streamlit as st
        for error in (KeyFormatError('PRIVATE-SECRET'), LedgerDataError('PRIVATE-SECRET'), SheetFormatError('PRIVATE-SECRET'), TimeoutError('PRIVATE-SECRET')):
            st.cache_resource.clear()
            with patch('storage.GoogleStore',side_effect=error):
                app=AppTest.from_file('app.py')
                app.secrets['login']={'salt':'test','password_hash':'test'}
                app.session_state['authenticated_until']=time.time()+1000
                app.run()
                self.assertFalse(app.exception)
                self.assertTrue(app.error)
                self.assertNotIn('PRIVATE-SECRET',str(app))

    def test_payoff_against_monthly_simulation(self):
        import random
        rng=random.Random(19)
        for _ in range(1000):
            remaining=rng.randrange(1,100_000_001); monthly=rng.randrange(1_000_000,10_000_001); paid=rng.randrange(0,20_000_001)
            left=remaining-max(monthly-paid,0); count=0
            while left>0: left-=monthly; count+=1
            self.assertEqual(payoff(remaining,monthly,paid,TODAY),month_add(TODAY,count))

    def test_fully_paid_and_interest_only(self):
        rows=[record(mother=100_000_000,me=100_000_000,total=200_000_000)]
        validate(rows[0],[],TODAY)
        validate(record(id='interest',mother=0,me=0,interest=100_000,total=100_000),rows,TODAY)
        self.assertEqual(balances(rows),{'mother':0,'me':0})
        self.assertEqual(payoff(0,0,0,TODAY),TODAY)

    def test_key_and_sheet_normalization(self):
        from cryptography.hazmat.primitives.asymmetric import rsa
        from cryptography.hazmat.primitives.serialization import Encoding, PrivateFormat, NoEncryption
        key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        pem = key.private_bytes(Encoding.PEM, PrivateFormat.PKCS8, NoEncryption()).decode()
        for value in (pem, pem.replace('\n', '\\n'), '\n' + pem + '  '):
            self.assertEqual(service_account_info({'private_key': value})['private_key'], pem)
        with self.assertRaises(KeyFormatError):
            service_account_info({'private_key': '-----BEGIN PRIVATE KEY-----\nbroken\n-----END PRIVATE KEY-----'})
        rows = [[cell + ' ' for cell in HEADER] + ['', ''], ['id','time','save','{}','','']]
        self.assertEqual(clean_rows(rows), [HEADER, ['id','time','save','{}']])
        self.assertEqual(clean_rows([HEADER + ['keep']])[0][-1], 'keep')
    def test_validation(self):
        r=record(); validate(r,[],TODAY)
        self.assertEqual(balances([r]),{'mother':96_000_000,'me':99_000_000})
        for invalid in [record(total=1),record(me=101_000_000,total=105_000_000),record(date='2026-10-02'),record(interest=-1)]:
            with self.assertRaises(ValueError): validate(invalid,[],TODAY)
        validate(record(total=5_100_000,interest=100_000),[],TODAY)
    def test_payoff(self):
        self.assertEqual(payoff(100_000_000,4_000_000,0,TODAY),date(2028,10,1))
        self.assertEqual(payoff(100_000_000,1_000_000,0,TODAY),date(2035,1,1))
        self.assertEqual(payoff(96_000_000,4_000_000,4_000_000,TODAY),date(2028,10,1))
        self.assertIsNone(payoff(1,0,0,TODAY))
        self.assertEqual(payoff(0,0,0,TODAY),TODAY)
        self.assertEqual(payoff(500_000,1_000_000,0,TODAY),TODAY)
        self.assertEqual(payoff(500_000,1_000_000,1_000_000,TODAY),date(2026,11,1))
        self.assertEqual(payoff(80_000_000,1_000_000,20_000_000,TODAY),date(2033,6,1))
    def test_excel(self):
        rows=[record(memo='=HYPERLINK("evil")')]
        wb=load_workbook(BytesIO(excel(rows,DEFAULT_PLAN,TODAY)))
        self.assertEqual(wb['상환 내역']['I2'].data_type,'s')
        self.assertEqual(wb['잔액 요약']['D2'].value,96_000_000)
    def test_events_and_conflict(self):
        s=fake_store(); rev=s.load()[2]
        s.mutate('save',record(),rev,TODAY)
        with self.assertRaises(ValueError): s.mutate('plan',DEFAULT_PLAN,rev,TODAY)
        s.mutate('save',record(total=6_000_000,me=2_000_000),s.load()[2],TODAY)
        self.assertEqual(len(s.load()[0]),1)
        s.mutate('delete',{'id':'a'},s.load()[2],TODAY)
        self.assertEqual(s.load()[0],[])
    def test_login_blocks(self):
        app=AppTest.from_file('app.py').run()
        self.assertFalse(app.exception)
        self.assertTrue(app.error)
        self.assertEqual(len(app.dataframe),0)
    def test_dashboard(self):
        s=fake_store()
        with patch('storage.GoogleStore',return_value=s):
            app=AppTest.from_file('app.py')
            app.secrets['login']={'salt':'test','password_hash':'test'}
            app.session_state['authenticated_until']=time.time()+1000
            app.run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.get('plotly_chart')),6)
            self.assertEqual(len(app.tabs),12)
            for b in app.button:
                if b.label=='확인한 금액 저장': b.click(); break
            app.run()
            self.assertFalse(app.exception)
            self.assertTrue(app.error)
            self.assertEqual(s.load()[0],[])
            for widget in app.text_input:
                if widget.label=='송금 은행': widget.set_value('국민은행')
                if widget.label=='실제 송금자명': widget.set_value('본인')
            for widget in app.number_input:
                if widget.label=='엄마 원금 상환액': widget.set_value(4_000_000)
                if widget.label=='본인 원금 상환액': widget.set_value(1_000_000)
                if widget.label=='실제 총 송금액': widget.set_value(5_000_000)
            for button in app.button:
                if button.label=='확인한 금액 저장': button.click(); break
            app.run()
            self.assertFalse(app.exception)
            self.assertEqual(len(s.load()[0]),1)
            self.assertEqual(len(app.dataframe),1)
            self.assertTrue(app.success)
if __name__=='__main__': unittest.main()
