from datetime import date
from calendar import monthrange
import hashlib
import json
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

PRINCIPAL = 100_000_000
FIELDS = ['id', 'date', 'bank', 'sender', 'total', 'mother', 'me', 'interest', 'memo']
DEFAULT_PLAN = {'mother': 4_000_000, 'me': 1_000_000}
MAX_AMOUNT = 2**53 - 1

def amount_words(amount):
    if type(amount) is not int or amount < 0 or amount > MAX_AMOUNT:
        raise ValueError('금액은 0 이상의 정수여야 합니다.')
    if amount == 0:
        return '영 원'
    digits = ['', '일', '이', '삼', '사', '오', '육', '칠', '팔', '구']
    groups = []
    for large in ('', '만', '억', '조'):
        amount, group = divmod(amount, 10000)
        text = ''
        for divisor, unit in ((1000, '천'), (100, '백'), (10, '십'), (1, '')):
            digit, group = divmod(group, divisor)
            if digit:
                text += ('' if digit == 1 and unit else digits[digit]) + unit
        if text:
            if text == '일' and large == '만':
                text = ''
            groups.append(text + large)
    return ' '.join(reversed(groups)) + ' 원'

def validate_plan(plan):
    if not isinstance(plan, dict) or set(plan) != {'mother', 'me'} or any(type(v) is not int or not 0 <= v <= MAX_AMOUNT for v in plan.values()):
        raise ValueError('월 상환 목표는 0 이상의 정수로 입력해주세요.')

def digest(rows):
    return hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode()).hexdigest()

def fold(events):
    records, plan = {}, dict(DEFAULT_PLAN)
    seen = set()
    for row in events:
        if len(row) != 4:
            raise ValueError('저장된 이벤트 형식이 올바르지 않습니다.')
        eid, timestamp, kind, raw = row
        if not eid or not timestamp:
            raise ValueError('저장된 이벤트 식별 정보가 없습니다.')
        if eid in seen:
            continue
        seen.add(eid)
        data = json.loads(raw)
        if not isinstance(data, dict):
            raise ValueError('저장된 이벤트 내용은 객체여야 합니다.')
        if kind == 'save':
            validate(data, [], date.max)
            records[data['id']] = data
        elif kind == 'delete':
            if not isinstance(data.get('id'), str) or not data['id']:
                raise ValueError('삭제할 내역 식별 정보가 없습니다.')
            records.pop(data['id'], None)
        elif kind == 'plan':
            validate_plan(data)
            plan = data
        elif kind in ('hana_save', 'hana_delete','gf_save','gf_delete'):
            # 별도 대출의 이벤트는 가족 장부의 잔액과 계획에 반영하지 않는다.
            pass
        elif kind == 'backup_confirm':
            from backup import validate_backup
            validate_backup(data,date.max)
        else:
            raise ValueError('알 수 없는 저장 이벤트입니다.')
    result = sorted(records.values(), key=lambda r: (r['date'], r['id']))
    if any(value < 0 for value in balances(result).values()):
        raise ValueError('저장된 원금 합계가 대출액을 초과합니다.')
    return result, plan

def balances(records):
    return {p: PRINCIPAL - sum(int(r[p]) for r in records) for p in ('mother', 'me')}

def validate(record, records, today):
    if not isinstance(record, dict) or any(field not in record for field in FIELDS):
        raise ValueError('상환 내역의 필수 항목이 없습니다.')
    for field in ('id', 'date', 'bank', 'sender', 'memo'):
        if not isinstance(record[field], str) or ILLEGAL_CHARACTERS_RE.search(record[field]):
            raise ValueError('입력한 글의 형식이나 특수문자를 확인해주세요.')
    if not record['id'].strip() or len(record['id']) > 100:
        raise ValueError('상환 내역 식별 정보가 올바르지 않습니다.')
    try:
        when = date.fromisoformat(record['date'])
        if when.isoformat() != record['date']:
            raise ValueError()
    except ValueError:
        raise ValueError('상환일은 YYYY-MM-DD 형식으로 입력해주세요.') from None
    if when > today:
        raise ValueError('미래 날짜는 실제 상환 내역으로 저장할 수 없습니다.')
    if not record['bank'].strip() or not record['sender'].strip():
        raise ValueError('은행과 실제 송금자명을 입력해주세요.')
    for k in ('total', 'mother', 'me', 'interest'):
        if type(record[k]) is not int or not 0 <= record[k] <= MAX_AMOUNT:
            raise ValueError('금액은 0 이상의 정수로 입력해주세요.')
    if record['total'] <= 0 or record['total'] != record['mother'] + record['me'] + record['interest']:
        raise ValueError('송금액은 엄마 원금 + 본인 원금 + 이자의 합계와 같아야 합니다.')
    others = [r for r in records if r['id'] != record['id']]
    remaining = balances(others)
    if any(record[p] > remaining[p] for p in remaining):
        raise ValueError('각자의 남은 원금보다 많이 상환할 수 없습니다.')
    if len(record['memo']) > 1000 or len(record['bank']) > 100 or len(record['sender']) > 100:
        raise ValueError('입력한 글이 너무 깁니다.')

def month_paid(records, today):
    return {p: sum(r[p] for r in records if r['date'][:7] == today.strftime('%Y-%m')) for p in ('mother', 'me')}

def month_add(today, count):
    m = today.year * 12 + today.month - 1 + count
    y, mo = divmod(m, 12)
    return date(y, mo + 1, min(today.day, monthrange(y, mo + 1)[1]))

def payoff(remaining, monthly, paid_this_month, today):
    if remaining <= 0:
        return today
    if monthly <= 0:
        return None
    current = max(monthly - paid_this_month, 0)
    later = max(remaining - current, 0)
    count = (later + monthly - 1) // monthly
    try:
        return month_add(today, count)
    except ValueError:
        return None

def excel(records, plan, today):
    wb = Workbook()
    ws = wb.active
    ws.title = '잔액 요약'
    ws.append(['구분', '빌린 원금', '갚은 원금', '남은 원금', '월 목표', '예상 완납월'])
    remaining, paid = balances(records), month_paid(records, today)
    for p, label in [('mother', '엄마'), ('me', '본인')]:
        d = payoff(remaining[p], plan[p], paid[p], today)
        ws.append([label, PRINCIPAL, PRINCIPAL - remaining[p], remaining[p], plan[p], d.strftime('%Y-%m') if d else '계산 불가'])
    ws.append(['전체', PRINCIPAL * 2, PRINCIPAL * 2 - sum(remaining.values()), sum(remaining.values()), sum(plan.values())])
    ws = wb.create_sheet('상환 내역')
    ws.append(['기록 ID', '상환일', '은행', '실제 송금자', '총 송금액', '엄마 원금', '본인 원금', '이자', '메모'])
    for r in records:
        ws.append([r[f] for f in FIELDS])
    ws = wb.create_sheet('월별 합계')
    ws.append(['월', '엄마 원금', '본인 원금', '이자', '총 송금액'])
    for month in sorted({r['date'][:7] for r in records}):
        subset = [r for r in records if r['date'][:7] == month]
        ws.append([month] + [sum(r[k] for r in subset) for k in ('mother', 'me', 'interest', 'total')])
    for ws in wb:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='2563EB')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(45, max(18, max(len(str(c.value or '')) for c in column) + 3))
            for cell in column[1:]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # 외부 입력을 Excel 수식으로 실행하지 않음
                if isinstance(cell.value, int):
                    cell.number_format = '#,##0'
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
